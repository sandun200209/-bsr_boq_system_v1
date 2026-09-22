from __future__ import annotations
import io
import openpyxl
import pytest
from app.database import SessionLocal
from app.models import (
    Project,
    ProjectSection,
    ProjectItem,
    DuplicationRecord,
    ChangeRegisterRecord,
    Template,
    TemplateMapping,
)
from app.services.project_service import (
    get_or_create_default_project,
    compare_historical_rates,
)
from app.services.template_manager_service import (
    get_or_register_default_template,
    export_project_excel,
)

@pytest.fixture(scope="module")
def db_session():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()

@pytest.fixture(scope="module")
def default_project(db_session):
    return get_or_create_default_project(db_session)


class TestProjectEstimatingAndTemplateExport:

    def test_separate_project_year_and_rate_source_year(self, db_session, default_project):
        """
        Rule: Project Base Year and Rate Source Year must be strictly separate fields.
        Item 1 can use 2026, Item 2 can use 2023, while Project Base Year is 2026.
        """
        assert default_project.project_base_year == 2026

        demo_sec = next(s for s in default_project.sections if s.section_code == "SEC-DEMO")
        items = demo_sec.items
        assert len(items) >= 2

        # Item 1 uses 2026
        item1 = next(it for it in items if it.item_no == "1")
        assert item1.rate_source_year == "2026"
        assert item1.rate_source_book == "BSR Southern 2026"

        # Item 2 uses 2023 (distinct from project base year 2026)
        item2 = next(it for it in items if it.item_no == "2")
        assert item2.rate_source_year == "2023"
        assert item2.rate_source_year != str(default_project.project_base_year)
        assert item2.adjustment == 15.0

    def test_historical_rate_comparison_side_by_side(self, db_session):
        """
        For a query or master item, returns all available historical years side by side.
        """
        comp = compare_historical_rates(db_session, query="brickwork", base_year=2026)
        assert comp.project_base_year == 2026
        assert len(comp.all_rates) > 0
        assert len(comp.rates_by_year) >= 1
        assert comp.average_rate > 0.0

    def test_explicit_user_rate_selection_not_overwritten(self, db_session, default_project):
        """
        User selected rate, year, and adjustment must not be automatically overwritten.
        """
        demo_sec = next(s for s in default_project.sections if s.section_code == "SEC-DEMO")
        item2 = next(it for it in demo_sec.items if it.item_no == "2")
        orig_rate = item2.selected_rate
        orig_year = item2.rate_source_year

        # Simulate editing project base year or other fields
        default_project.project_name = "DGH Matara - Updated Title"
        db_session.commit()
        db_session.refresh(item2)

        assert item2.selected_rate == orig_rate
        assert item2.rate_source_year == orig_year

    def test_independent_project_sections(self, default_project):
        """
        Each output section maintains its own dataset and sort order independently.
        """
        section_codes = [s.section_code for s in default_project.sections]
        assert "SEC-DEMO" in section_codes
        assert "SEC-SITE-DUP" in section_codes
        assert "SEC-REV7-REG" in section_codes
        assert "SEC-ELEC-REC" in section_codes

        # Sections have different types and independent row collections
        sec_types = {s.section_code: s.section_type for s in default_project.sections}
        assert sec_types["SEC-DEMO"] == "BOQ"
        assert sec_types["SEC-SITE-DUP"] == "DUPLICATION"
        assert sec_types["SEC-REV7-REG"] == "CHANGE_REGISTER"
        assert sec_types["SEC-ELEC-REC"] == "RECONCILIATION"

    def test_ot_site_works_duplication_export(self, default_project):
        """
        Exports 'MATARA OT RENOVATION – SITE WORKS BOQ REVIEWED FOR DUPLICATION'
        with exact BOQ-style columns:
        Item, BSR Ref, Description, Unit, Qty, Rate, Amount, Rate Source Year, Duplicate With, Status, Remarks.
        """
        site_sec = next(s for s in default_project.sections if "site works" in s.section_name.lower())
        excel_buf = export_project_excel(default_project, mode="separate", target_section_id=site_sec.id)

        wb = openpyxl.load_workbook(excel_buf, data_only=False)
        assert "Site Works - Reviewed" in wb.sheetnames
        ws = wb["Site Works - Reviewed"]

        # Check title
        assert "SITE WORKS BOQ REVIEWED FOR DUPLICATION" in str(ws.cell(1, 1).value)

        # Check row 5 data
        row5 = [ws.cell(5, c).value for c in range(1, 12)]
        assert row5[0] == "1"  # Item
        assert row5[1] == "A. Preliminaries"  # BSR Ref
        assert "Protection to live hospital" in str(row5[2])  # Description
        assert row5[3] == "Item"  # Unit
        assert row5[4] == 1.0  # Qty
        assert row5[5] == 250000.0  # Rate
        assert row5[6] == "=E5*F5"  # Amount live formula
        assert row5[7] == "2026"  # Rate Source Year
        assert row5[8] == "None"  # Duplicate With
        assert row5[9] == "RETAIN"  # Status

    def test_rev7_change_register_export(self, default_project):
        """
        Exports 'REV 7 CONSOLIDATION CHANGE / DUPLICATION REGISTER' with exact columns:
        Ref, Package / Sheet, Item, Description, Original Status, REV 7 Action, Duplicate With, Rate Source, Cost Impact, Reason.
        """
        reg_sec = next(s for s in default_project.sections if "change" in s.section_name.lower())
        excel_buf = export_project_excel(default_project, mode="separate", target_section_id=reg_sec.id)

        wb = openpyxl.load_workbook(excel_buf, data_only=False)
        assert "Consolidation Change Register" in wb.sheetnames
        ws = wb["Consolidation Change Register"]

        assert "REV 7 CONSOLIDATION CHANGE / DUPLICATION REGISTER" in str(ws.cell(1, 1).value)

        # Check row 4 data
        row4 = [ws.cell(4, c).value for c in range(1, 11)]
        assert row4[0] == "CR-01"  # Ref
        assert row4[1] == "Electrical"  # Package / Sheet
        assert row4[2] == "E-01"  # Item
        assert "40 OT light points" in str(row4[3])  # Description
        assert "Source subtotal" in str(row4[4])  # Original Status
        assert "Reduce electrical quantities" in str(row4[5])  # REV 7 Action
        assert "Modular OT C6" in str(row4[6])  # Duplicate With
        assert row4[7] == "BSR 2026"  # Rate Source
        assert row4[8] == 1080000.0  # Cost Impact
        assert "Modular OT package" in str(row4[9])  # Reason

    def test_datasets_not_combined(self, default_project):
        """
        Rule: Do NOT combine the OT Site Works Duplication and REV 7 Change Register datasets.
        """
        site_sec = next(s for s in default_project.sections if "site works" in s.section_name.lower())
        reg_sec = next(s for s in default_project.sections if "change" in s.section_name.lower())

        # Datasets are distinct models and separate section IDs
        assert site_sec.id != reg_sec.id
        assert len(site_sec.duplication_records) > 0
        assert len(site_sec.change_register_records) == 0

        assert len(reg_sec.change_register_records) > 0
        assert len(reg_sec.duplication_records) == 0

    def test_template_manager_mapping(self, db_session):
        """
        Template manager allows master Excel workbook to be registered and mapped dynamically.
        """
        tpl = get_or_register_default_template(db_session)
        assert tpl.template_name is not None
        assert len(tpl.mappings) >= 5

        # Check mapping fields exist
        m = tpl.mappings[0]
        assert hasattr(m, "target_sheet")
        assert hasattr(m, "start_row")
        assert hasattr(m, "item_no_column")
        assert hasattr(m, "description_column")
        assert hasattr(m, "rate_column")

    def test_export_options_consolidated_and_separate(self, default_project):
        """
        Export Option A: One Excel workbook with each section as a separate worksheet.
        Export Option B: Separate Excel files for each selected section.
        """
        # Option A: Consolidated
        excel_cons = export_project_excel(default_project, mode="consolidated")
        wb_cons = openpyxl.load_workbook(excel_cons)
        assert len(wb_cons.sheetnames) >= 4
        assert "Site Works - Reviewed" in wb_cons.sheetnames
        assert "Consolidation Change Register" in wb_cons.sheetnames
        assert "Electrical BOQ - Reviewed" in wb_cons.sheetnames

        # Option B: Separate
        site_sec = next(s for s in default_project.sections if "site works" in s.section_name.lower())
        excel_sep = export_project_excel(default_project, mode="separate", target_section_id=site_sec.id)
        wb_sep = openpyxl.load_workbook(excel_sep)
        assert len(wb_sep.sheetnames) == 1
        assert wb_sep.sheetnames[0] == "Site Works - Reviewed"
